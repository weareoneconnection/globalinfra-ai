#!/usr/bin/env python3
"""Build the AI Outbound competition design report from the official template."""

from __future__ import annotations

import re
import shutil
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "submission" / "report_source.md"
TEMPLATE = Path("/Users/maqing/Downloads/附件3：AI出海赛道设计报告模板.docx")
OUTPUT_DIR = ROOT / "submission" / "final"
ASSET_DIR = ROOT / "submission" / "assets"
OUTPUT_DOCX = OUTPUT_DIR / "GlobalInfra_AI_AI出海应用设计报告.docx"
SCREENSHOT = Path("/private/tmp/globalinfra-phase91-production.png")

FONT_REGULAR = Path("/System/Library/AssetsV2/com_apple_MobileAsset_Font8/1b6b2f148cdf17dd0d788706c83c61ca1998c3a6.asset/AssetData/STFANGSO.ttf")
FONT_UI = Path("/System/Library/AssetsV2/com_apple_MobileAsset_Font8/86ba2c91f017a3749571a82f2c6d890ac7ffb2fb.asset/AssetData/PingFang.ttc")
if not FONT_UI.exists():
    FONT_UI = Path("/System/Library/Fonts/PingFang.ttc")

NAVY = "17365D"
BLUE = "2F75B5"
PALE_BLUE = "EAF2F8"
GRID = "D9E1F2"
TEXT = RGBColor(0, 0, 0)


def pil_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    path = FONT_UI if FONT_UI.exists() else FONT_REGULAR
    index = 1 if bold and path.suffix.lower() == ".ttc" else 0
    try:
        return ImageFont.truetype(str(path), size=size, index=index)
    except OSError:
        return ImageFont.truetype(str(FONT_REGULAR), size=size)


def rounded_box(draw: ImageDraw.ImageDraw, xy: tuple[int, int, int, int], fill: str, outline: str = "#7EA6C9", radius: int = 22, width: int = 3) -> None:
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def centered_multiline(draw: ImageDraw.ImageDraw, xy: tuple[int, int, int, int], text: str, font: ImageFont.FreeTypeFont, fill: str = "#102A43", spacing: int = 8) -> None:
    x1, y1, x2, y2 = xy
    box = draw.multiline_textbbox((0, 0), text, font=font, spacing=spacing, align="center")
    w, h = box[2] - box[0], box[3] - box[1]
    draw.multiline_text(((x1 + x2 - w) / 2, (y1 + y2 - h) / 2), text, font=font, fill=fill, spacing=spacing, align="center")


def create_architecture(path: Path) -> None:
    img = Image.new("RGB", (1800, 1080), "white")
    draw = ImageDraw.Draw(img)
    title_font = pil_font(48, True)
    label_font = pil_font(31, True)
    text_font = pil_font(24)
    draw.text((900, 45), "GlobalInfra AI 总体架构", font=title_font, fill="#0B1F33", anchor="ma")

    rows = [
        (130, "接入层", ["文本与文档", "现场 PWA", "企业连接器", "BIM GIS", "IoT 遥测"], "#EAF2F8"),
        (315, "智能层", ["结构化接入", "国家包检索", "多智能体", "模型路由", "策略与拒答"], "#D9EAF7"),
        (500, "项目数字主线", ["项目与阶段", "事件与工单", "主张与证据", "决策与行动", "核验与经验"], "#CBE3F6"),
        (685, "数字孪生层", ["Twin ID", "IFC GLTF", "3D Tiles GIS", "4D 5D", "调试与运营"], "#DDEBF7"),
        (870, "治理与部署", ["身份与 RBAC", "人工审批", "执行回执", "独立核验", "区域与边缘"], "#E8EEF5"),
    ]
    for row_index, (y, label, items, fill) in enumerate(rows):
        rounded_box(draw, (80, y, 300, y + 130), "#17365D", "#17365D", radius=18)
        centered_multiline(draw, (80, y, 300, y + 130), label, label_font, "white")
        box_w, gap = 252, 18
        start_x = 340
        for i, item in enumerate(items):
            x = start_x + i * (box_w + gap)
            rounded_box(draw, (x, y, x + box_w, y + 130), fill, "#7EA6C9", radius=18)
            centered_multiline(draw, (x + 12, y + 8, x + box_w - 12, y + 122), item, text_font)
        if row_index < len(rows) - 1:
            draw.line((900, y + 132, 900, y + 177), fill="#2F75B5", width=6)
            draw.polygon([(884, y + 164), (916, y + 164), (900, y + 184)], fill="#2F75B5")
    img.save(path, quality=95)


def create_golden_thread(path: Path) -> None:
    img = Image.new("RGB", (2000, 720), "white")
    draw = ImageDraw.Draw(img)
    title_font = pil_font(48, True)
    number_font = pil_font(28, True)
    text_font = pil_font(24, True)
    small_font = pil_font(19)
    draw.text((1000, 55), "马来西亚项目七步黄金主线", font=title_font, fill="#0B1F33", anchor="ma")
    steps = [
        ("1", "双语事件\n接入", "HSE\nCRITICAL"),
        ("2", "国家与阶段\n路由", "MY\n施工"),
        ("3", "国家包\n应用", "DOSH\nCIDB"),
        ("4", "知识证据\n检索", "来源\n版本"),
        ("5", "数字孪生\n定位", "Twin ID\n4D 5D"),
        ("6", "智能体与\n人工审批", "7 Agents\n2 Approvals"),
        ("7", "执行与\n独立核验", "Receipt\nVerified"),
    ]
    w, h, gap = 230, 360, 44
    start_x, y = 80, 185
    for i, (num, label, proof) in enumerate(steps):
        x = start_x + i * (w + gap)
        fill = "#EAF2F8" if i < 5 else "#D8EAD8"
        outline = "#2F75B5" if i < 5 else "#3D7C47"
        rounded_box(draw, (x, y, x + w, y + h), fill, outline, radius=24, width=4)
        draw.ellipse((x + 82, y - 38, x + 148, y + 28), fill="#17365D", outline="#17365D")
        draw.text((x + 115, y - 5), num, font=number_font, fill="white", anchor="mm")
        centered_multiline(draw, (x + 16, y + 38, x + w - 16, y + 205), label, text_font)
        centered_multiline(draw, (x + 16, y + 215, x + w - 16, y + 330), proof, small_font, "#35546F")
        if i < len(steps) - 1:
            ax = x + w + 7
            draw.line((ax, y + h / 2, ax + gap - 14, y + h / 2), fill="#6889A6", width=6)
            draw.polygon([(ax + gap - 30, y + h / 2 - 14), (ax + gap - 30, y + h / 2 + 14), (ax + gap - 10, y + h / 2)], fill="#6889A6")
    draw.text((1000, 625), "AI 建议  →  人工批准  →  工具执行  →  独立核验  →  可审计关闭", font=pil_font(29, True), fill="#17365D", anchor="ma")
    img.save(path, quality=95)


def set_run_font(run, size: float, bold: bool = False, color: RGBColor = TEXT, name: str = "PingFang SC") -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:hint"), "eastAsia" if name == "PingFang SC" else "default")
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:cs"), name)


def add_inline_runs(paragraph, text: str, size: float, bold: bool = False) -> None:
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        code = part.startswith("`") and part.endswith("`")
        strong = part.startswith("**") and part.endswith("**")
        content = part[1:-1] if code else part[2:-2] if strong else part
        run = paragraph.add_run(content)
        set_run_font(run, 10.5 if code else size, bold=bold or strong, name="Consolas" if code else "PingFang SC")


def add_title(doc: Document, text: str):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(108)
    p.paragraph_format.space_after = Pt(36)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    set_run_font(run, 22, bold=True)
    return p


def add_cover_line(doc: Document, text: str, index: int):
    p = doc.add_paragraph()
    boundary = text.startswith("数据边界")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if boundary else WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0 if boundary else 6)
    if index == 0:
        size, bold = 16, True
    elif boundary:
        size, bold = 12, False
        p.paragraph_format.first_line_indent = Inches(0.44)
        p.paragraph_format.space_before = Pt(36)
    else:
        size, bold = 14, False
    add_inline_runs(p, text, size=size, bold=bold)
    return p


def add_heading(doc: Document, text: str, level: int):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.widow_control = True
    if level == 1:
        p.paragraph_format.page_break_before = True
        p.paragraph_format.first_line_indent = Inches(0.446)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.5
        size = 16
    elif level == 2:
        p.paragraph_format.first_line_indent = Inches(0.446)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.35
        size = 15
    else:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.25
        size = 14
    run = p.add_run(text)
    set_run_font(run, size, bold=True)
    return p


def add_body(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(0.444)
    p.paragraph_format.line_spacing = 1.45
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.widow_control = True
    add_inline_runs(p, text, 15)
    return p


def add_bullet(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.58)
    p.paragraph_format.first_line_indent = Inches(-0.28)
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("• ")
    set_run_font(run, 14, bold=True)
    add_inline_runs(p, text, 14)
    return p


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = GRID, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_pr = table._tbl.tblPr
    tbl_width = table_pr.find(qn("w:tblW"))
    if tbl_width is not None:
        tbl_width.set(qn("w:type"), "pct")
        tbl_width.set(qn("w:w"), "5000")

    max_chars = [max(max(len(cell), 4) for cell in (row[i] for row in rows)) for i in range(cols)]
    total = sum(min(v, 40) for v in max_chars)
    ratios = [max(0.14, min(0.46, min(v, 40) / total)) for v in max_chars]
    ratio_total = sum(ratios)
    widths = [6.0 * r / ratio_total for r in ratios]

    for r_idx, source_row in enumerate(rows):
        tr_pr = table.rows[r_idx]._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        for c_idx, value in enumerate(source_row):
            cell = table.cell(r_idx, c_idx)
            cell.width = Inches(widths[c_idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_borders(cell)
            set_cell_shading(cell, NAVY if r_idx == 0 else (PALE_BLUE if r_idx % 2 == 0 else "FFFFFF"))
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (r_idx == 0 or len(value) <= 12) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run(value)
            set_run_font(run, 10.5, bold=r_idx == 0, color=RGBColor(255, 255, 255) if r_idx == 0 else TEXT)
    header_tr = table.rows[0]._tr
    tr_pr = header_tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_figure(doc: Document, image_path: Path, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    if image_path.name == "architecture.png":
        figure_width = 5.25
    elif image_path.name == "golden-thread.png":
        figure_width = 5.0
    elif image_path.name.startswith("assistant-"):
        # Tall UI screenshots: narrower so two fit on one page.
        figure_width = 4.3
    else:
        figure_width = 5.85
    picture = run.add_picture(str(image_path), width=Inches(figure_width))
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", caption)
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.keep_with_next = False
    cp.paragraph_format.space_after = Pt(8)
    cr = cp.add_run(caption)
    set_run_font(cr, 12)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        values = [value.strip() for value in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", value.replace(" ", "")) for value in values):
            rows.append(values)
        index += 1
    return rows, index


def clear_document_body(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def remove_embedded_subset_fonts(docx_path: Path) -> None:
    """Remove template-only font subsets that omit glyphs used by the full report."""
    temp_path = docx_path.with_suffix(".fonts-fixed.docx")
    with zipfile.ZipFile(docx_path, "r") as source_zip, zipfile.ZipFile(
        temp_path, "w", compression=zipfile.ZIP_DEFLATED
    ) as target_zip:
        for info in source_zip.infolist():
            if info.filename.startswith("word/fonts/"):
                continue
            data = source_zip.read(info.filename)
            if info.filename == "word/fontTable.xml":
                data = re.sub(rb"<w:embed(?:Regular|Bold|Italic|BoldItalic)\b[^>]*/>", b"", data)
            elif info.filename == "word/_rels/fontTable.xml.rels":
                data = re.sub(
                    rb"<Relationship\b[^>]*relationships/font[^>]*/>",
                    b"",
                    data,
                )
            elif info.filename.startswith("word/footer") and info.filename.endswith(".xml"):
                def add_footer_alt(match: re.Match[bytes]) -> bytes:
                    attributes = match.group(1)
                    if b"descr=" not in attributes:
                        attributes += b' descr="Page number text box"'
                    if b"title=" not in attributes:
                        attributes += b' title="Page number"'
                    return b"<wp:docPr" + attributes + b"/>"

                data = re.sub(rb"<wp:docPr\b([^>]*)/>", add_footer_alt, data)
            target_zip.writestr(info, data)
    temp_path.replace(docx_path)


def build_report() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    create_architecture(ASSET_DIR / "architecture.png")
    create_golden_thread(ASSET_DIR / "golden-thread.png")
    if SCREENSHOT.exists():
        shutil.copy2(SCREENSHOT, ASSET_DIR / "phase91-production.png")
    else:
        raise FileNotFoundError(f"Production verification screenshot is missing: {SCREENSHOT}")

    doc = Document(TEMPLATE)
    clear_document_body(doc)
    if "Title" not in [style.name for style in doc.styles]:
        doc.styles.add_style("Title", WD_STYLE_TYPE.PARAGRAPH)
    doc.core_properties.title = "GlobalInfra AI AI出海应用设计报告"
    doc.core_properties.subject = "第二届交融杯人工智能场景创新应用大赛 AI出海赛道"
    doc.core_properties.author = "GlobalInfra AI 参赛团队"
    doc.core_properties.keywords = "AI出海,交通基础设施,数字孪生,多智能体,可信执行"

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    first_title = True
    cover_index = 0
    seen_heading = False
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line == "<PAGEBREAK>":
            p = doc.add_paragraph()
            p.add_run().add_break(WD_BREAK.PAGE)
            index += 1
            continue
        image_match = re.fullmatch(r"!\[(.+)]\((.+)\)", line)
        if image_match:
            caption, rel_path = image_match.groups()
            add_figure(doc, ROOT / rel_path, caption)
            index += 1
            continue
        if line.startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(doc, rows)
            continue
        if line.startswith("# ") and first_title:
            add_title(doc, line[2:].strip())
            first_title = False
            index += 1
            continue
        if line.startswith("## "):
            add_heading(doc, line[3:].strip(), 1)
            seen_heading = True
            index += 1
            continue
        if line.startswith("### "):
            add_heading(doc, line[4:].strip(), 2)
            index += 1
            continue
        if line.startswith("#### "):
            add_heading(doc, line[5:].strip(), 3)
            index += 1
            continue
        if line.startswith("- "):
            add_bullet(doc, line[2:].strip())
            index += 1
            continue
        if not seen_heading:
            add_cover_line(doc, line, cover_index)
            cover_index += 1
        else:
            add_body(doc, line)
        index += 1

    for section in doc.sections:
        section.start_type = WD_SECTION.NEW_PAGE
        section.page_width = Inches(8.2680556)
        section.page_height = Inches(11.6930556)
        section.left_margin = Inches(1.1020833)
        section.right_margin = Inches(1.1340278)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.header_distance = Inches(0.5909722)
        section.footer_distance = Inches(0.6888889)

    doc.save(OUTPUT_DOCX)
    remove_embedded_subset_fonts(OUTPUT_DOCX)
    return OUTPUT_DOCX


if __name__ == "__main__":
    print(build_report())
