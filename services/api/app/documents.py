"""Text extraction for universal intake documents (M1 scope: PDF / DOCX / TXT).

M1 deliberately covers text extraction only. Image OCR and voice ASR are
explicitly out of scope and deferred to a later version -- an unsupported
upload raises ``UnsupportedDocument`` so the API can return a clean 400 rather
than silently producing an empty case.

The PDF/DOCX parsers are imported lazily so the rest of the app (and the
offline ``create_all``/eval paths) never depends on those libraries being
importable just to start up.
"""
from __future__ import annotations

import io
import zipfile


class UnsupportedDocument(ValueError):
    """Raised for an upload outside the M1 text-extraction scope."""


_EXT_BY_CONTENT_TYPE = {
    "text/plain": "txt",
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/msword": "docx",
}

SUPPORTED_TYPES = ("txt", "pdf", "docx")
MAX_EXTRACTED_CHARACTERS = 500_000
MAX_DOCX_ENTRIES = 5_000
MAX_DOCX_UNCOMPRESSED_BYTES = 25 * 1024 * 1024


def detect_file_type(filename: str | None, content_type: str | None) -> str:
    """Resolve a normalized file type from filename extension first, then the
    declared content type. Raises ``UnsupportedDocument`` for anything else."""
    ext = ""
    if filename and "." in filename:
        ext = filename.rsplit(".", 1)[-1].lower()
    if ext in SUPPORTED_TYPES:
        return ext
    mapped = _EXT_BY_CONTENT_TYPE.get((content_type or "").split(";")[0].strip().lower())
    if mapped:
        return mapped
    raise UnsupportedDocument(
        f"Unsupported document type for '{filename}' (content-type {content_type!r}); "
        f"M1 supports only {', '.join(SUPPORTED_TYPES)} text extraction."
    )


def _extract_txt(content: bytes) -> str:
    if b"\x00" in content:
        raise UnsupportedDocument("TXT upload contains binary data.")
    return content.decode("utf-8", errors="replace")


def _extract_pdf(content: bytes) -> str:
    if not content.startswith(b"%PDF-"):
        raise UnsupportedDocument("File content is not a valid PDF.")
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise UnsupportedDocument("PDF extraction requires the 'pypdf' package.") from exc
    try:
        reader = PdfReader(io.BytesIO(content))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception as exc:
        raise UnsupportedDocument("PDF could not be parsed safely.") from exc
    return text[:MAX_EXTRACTED_CHARACTERS].strip()


def _extract_docx(content: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_DOCX_ENTRIES or sum(entry.file_size for entry in entries) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise UnsupportedDocument("DOCX expanded content exceeds safe parsing limits.")
            names = {entry.filename for entry in entries}
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise UnsupportedDocument("File content is not a valid DOCX document.")
    except zipfile.BadZipFile as exc:
        raise UnsupportedDocument("File content is not a valid DOCX document.") from exc
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise UnsupportedDocument("DOCX extraction requires the 'python-docx' package.") from exc
    try:
        document = Document(io.BytesIO(content))
        text = "\n".join(p.text for p in document.paragraphs)
    except Exception as exc:
        raise UnsupportedDocument("DOCX could not be parsed safely.") from exc
    return text[:MAX_EXTRACTED_CHARACTERS].strip()


def extract_text(filename: str | None, content: bytes, content_type: str | None = None) -> tuple[str, str]:
    """Extract plain text from an uploaded document.

    Returns ``(text, file_type)``. Raises ``UnsupportedDocument`` for a type
    outside the M1 text-extraction scope.
    """
    file_type = detect_file_type(filename, content_type)
    if file_type == "txt":
        return _extract_txt(content), file_type
    if file_type == "pdf":
        return _extract_pdf(content), file_type
    return _extract_docx(content), file_type
