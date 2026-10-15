"""Tests for the V1.3 M1 Universal Intake Foundation.

Covers the deterministic intake engine (Layer 1 wrapping classify + enrichment)
and both intake endpoints end to end against the real FastAPI app + seeded
dataset (see conftest.py). Every case here uses *new, non-seed* letter text, so
these exercise the "field input -> structured case -> specialist routing ->
analysis" path rather than the pre-seeded demo cases.
"""
import io
import pytest

from app.documents import detect_file_type, extract_text, UnsupportedDocument
from app.intake import build_intake_result, detect_language
from app.workflow import SPECIALISTS_BY_CASE_TYPE

# A brand-new owner letter, not present in the seed dataset.
NEW_OWNER_LETTER = (
    "Owner notice to the Contractor: the temporary access road serving the "
    "Malaysia viaduct has been damaged and earthwork encroachment was observed. "
    "Rectification is required within 48 hours. Refer to NCR-2026-311 and "
    "Clause C-14.2 of the contract."
)
NEW_HSE_LETTER_ZH = "现场巡检发现：工人在高处作业面靠近无防护边缘，缺少坠落防护，须立即处理。"


# --- unit: deterministic intake engine --------------------------------------

def test_layer1_case_type_and_risk_come_from_classifier():
    result = build_intake_result(NEW_OWNER_LETTER)
    # Layer 1 must agree with the classify.py contract.
    assert result["case_type"] == "Contract / Technical Interface"
    assert result["risk"] == "HIGH"


def test_recommended_agents_match_specialist_routing():
    result = build_intake_result(NEW_OWNER_LETTER)
    assert result["recommended_agents"] == SPECIALISTS_BY_CASE_TYPE["Contract / Technical Interface"]


def test_hse_letter_routes_to_hse_team_and_detects_chinese():
    result = build_intake_result(NEW_HSE_LETTER_ZH)
    assert result["case_type"] == "HSE"
    assert result["risk"] == "CRITICAL"
    assert result["recommended_agents"] == SPECIALISTS_BY_CASE_TYPE["HSE"]
    assert result["language"] == "zh"


def test_enrichment_fields_are_bilingual_and_populated():
    result = build_intake_result(NEW_OWNER_LETTER)
    assert result["case_type_label"] == {"en": "Contract / Technical Interface", "zh": "合同 / 技术接口"}
    assert result["risk_label"]["en"] and result["risk_label"]["zh"]
    assert all(set(i.keys()) == {"en", "zh"} for i in result["issues"])
    assert result["country"] == "MY"
    assert result["deadline"] == {"en": "48 hours", "zh": "48 小时"}
    assert "NCR-2026-311" in result["entities"]
    assert 0.0 < result["confidence"] <= 0.95
    assert result["required_approvals"]


def test_detect_language():
    assert detect_language("hello world") == "en"
    assert detect_language("你好，世界") == "zh"


# --- unit: document extraction ----------------------------------------------

def test_detect_file_type_from_extension_and_content_type():
    assert detect_file_type("a.txt", None) == "txt"
    assert detect_file_type("a.pdf", None) == "pdf"
    assert detect_file_type("a.docx", None) == "docx"
    assert detect_file_type("noext", "application/pdf") == "pdf"


def test_unsupported_document_rejected():
    import pytest
    with pytest.raises(UnsupportedDocument):
        detect_file_type("photo.png", "image/png")


def test_extract_txt_and_docx():
    text, ftype = extract_text("note.txt", "hello crack 裂缝".encode("utf-8"), "text/plain")
    assert ftype == "txt" and "裂缝" in text

    from docx import Document
    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Concrete surface honeycomb observed after formwork removal.")
    doc.save(buf)
    text, ftype = extract_text("obs.docx", buf.getvalue(), None)
    assert ftype == "docx" and "honeycomb" in text


# --- API: /intake/text ------------------------------------------------------

def _seeded_ids(client):
    return {c["id"] for c in client.get("/cases").json()}


def test_intake_text_creates_new_case_and_routes_to_specialists(client):
    before = _seeded_ids(client)
    r = client.post("/intake/text", json={"text": NEW_OWNER_LETTER})
    assert r.status_code == 200
    body = r.json()

    case_id = body["case_id"]
    assert case_id and case_id not in before  # a genuinely new, non-seed case
    assert body["intake"]["case_type"] == "Contract / Technical Interface"
    team_ids = [t["id"] for t in body["specialist_team"]]
    assert team_ids == SPECIALISTS_BY_CASE_TYPE["Contract / Technical Interface"]
    assert body["analysis"]["status"] == "AWAITING_APPROVAL"

    # The case really exists and the analysis workflow actually ran on it.
    detail = client.get(f"/cases/{case_id}").json()
    assert detail["status"] == "AWAITING_APPROVAL"
    assert len(detail["agent_runs"]) > 0
    assert len(detail["evidence"]) > 0
    # Routed to the Contract specialist (English label from /agents).
    assert any("Contract Agent" in run["agent"] for run in detail["agent_runs"])


def test_intake_text_hse_routes_to_hse_team(client):
    r = client.post("/intake/text", json={"text": NEW_HSE_LETTER_ZH})
    assert r.status_code == 200
    body = r.json()
    assert body["intake"]["case_type"] == "HSE"
    assert [t["id"] for t in body["specialist_team"]] == SPECIALISTS_BY_CASE_TYPE["HSE"]
    detail = client.get(f"/cases/{body['case_id']}").json()
    assert any("HSE" in run["agent"] for run in detail["agent_runs"])


def test_unrouted_intake_reaches_approval_gate_after_project_confirmation(client):
    text = "P47 pier: workers at height without edge protection, rectify within 48 hours."
    body = client.post("/intake/text", json={"text": text}).json()
    case_id = body["case_id"]
    assert body["intake"]["project_routing_status"] == "NEEDS_PROJECT_CONFIRMATION"
    assert client.get(f"/cases/{case_id}").json()["action"] is None

    r = client.post(f"/cases/{case_id}/confirm-project", json={"project_id": 1})
    assert r.status_code == 200
    assert r.json()["status"] == "AWAITING_APPROVAL"
    detail = client.get(f"/cases/{case_id}").json()
    assert detail["routing_status"] == "ROUTED"
    assert detail["action"]["status"] == "AWAITING_APPROVAL"
    assert any(a["case_id"] == case_id for a in client.get("/approvals").json())

    # Confirmation is a one-time human step, and the project must exist.
    assert client.post(f"/cases/{case_id}/confirm-project", json={"project_id": 1}).status_code == 409
    other = client.post("/intake/text", json={"text": text}).json()["case_id"]
    assert client.post(f"/cases/{other}/confirm-project", json={"project_id": 9999}).status_code == 404
    assert client.post("/cases/CASE-MISSING/confirm-project", json={"project_id": 1}).status_code == 404


def test_intake_text_is_bilingual_via_lang_param(client):
    en = client.post("/intake/text?lang=en", json={"text": NEW_OWNER_LETTER}).json()
    zh = client.post("/intake/text?lang=zh", json={"text": NEW_OWNER_LETTER}).json()
    assert en["intake"]["case_type_label"] == "Contract / Technical Interface"
    assert zh["intake"]["case_type_label"] == "合同 / 技术接口"
    assert en["intake"]["risk_label"] == "High"
    assert zh["intake"]["risk_label"] == "高"


def test_intake_text_rejects_empty(client):
    r = client.post("/intake/text", json={"text": "   "})
    assert r.status_code == 422


def test_intake_dispatches_classification_and_extraction_concurrently(client, monkeypatch):
    """M6: the two independent LLM layers are scheduled concurrently via
    asyncio.gather, not awaited one after another.

    Each layer is stubbed to record an 'enter' marker, yield control once, then
    record an 'exit' marker. Under concurrent dispatch both coroutines reach
    'enter' before either reaches 'exit' (extract_enter precedes classify_exit);
    serial ``await`` would fully finish classify before extract even starts.
    """
    import asyncio
    import app.main as main

    order: list[str] = []
    real_classify = main.classify_text_llm

    async def classify_stub(text, *args, **kwargs):
        order.append("classify_enter")
        await asyncio.sleep(0)
        order.append("classify_exit")
        return await real_classify(text, *args, **kwargs)  # keep the real rules result

    async def extract_stub(text, *args, **kwargs):
        order.append("extract_enter")
        await asyncio.sleep(0)
        order.append("extract_exit")
        return None  # deterministic extraction path, unchanged downstream

    monkeypatch.setattr(main, "classify_text_llm", classify_stub)
    monkeypatch.setattr(main, "extract_fields_llm", extract_stub)

    r = client.post("/intake/text", json={"text": NEW_OWNER_LETTER})
    assert r.status_code == 200
    assert r.json()["intake"]["case_type"] == "Contract / Technical Interface"
    # Concurrent: extract started before classify finished. Serial would fail here.
    assert order.index("extract_enter") < order.index("classify_exit")


# --- API: /intake/document --------------------------------------------------

def test_intake_document_txt_creates_case(client):
    content = NEW_OWNER_LETTER.encode("utf-8")
    r = client.post("/intake/document", files={"file": ("letter.txt", content, "text/plain")})
    assert r.status_code == 200
    body = r.json()
    assert body["intake"]["case_type"] == "Contract / Technical Interface"
    assert body["submission_id"]
    detail = client.get(f"/cases/{body['case_id']}").json()
    assert detail["status"] == "AWAITING_APPROVAL"


def test_intake_document_docx_creates_quality_case(client):
    from docx import Document
    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Quality observation: concrete surface honeycomb found at the drainage wall after formwork removal.")
    doc.save(buf)
    r = client.post(
        "/intake/document",
        files={"file": ("obs.docx", buf.getvalue(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert r.status_code == 200
    body = r.json()
    assert body["intake"]["case_type"] == "Quality"
    assert [t["id"] for t in body["specialist_team"]] == SPECIALISTS_BY_CASE_TYPE["Quality"]


def test_intake_document_image_only_pdf_is_preserved_and_queued_for_ocr(client):
    from pypdf import PdfWriter

    buf = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=200)
    writer.write(buf)
    content = buf.getvalue()

    r = client.post(
        "/intake/document?lang=en&data_classification=internal",
        files={"file": ("scanned-letter.pdf", content, "application/pdf")},
    )
    assert r.status_code == 200
    body = r.json()
    assert body["case_id"] is None
    assert body["analysis"] is None
    assert body["specialist_team"] == []
    assert body["document"]["parse_status"] == "OCR_REQUIRED"
    assert body["document"]["byte_size"] == len(content)
    assert body["document"]["original_preserved"] is True
    receipt = body["ingestion_receipt"]
    assert receipt["status"] == "OCR_REQUIRED"
    assert receipt["target_type"] == "SourceDocument"
    assert receipt["validation_summary"]["ocr_required"] is True
    assert receipt["lineage"]["raw_preserved"] is True

    from app.database import SessionLocal
    from app.models import SourceDocument

    with SessionLocal() as db:
        stored = db.get(SourceDocument, body["document"]["id"])
        assert stored is not None
        assert stored.original_content == content
        assert stored.storage_ref == f"sha256:{receipt['checksum']}"


def test_intake_document_rejects_unsupported_type(client):
    r = client.post("/intake/document", files={"file": ("photo.png", b"\x89PNG\r\n", "image/png")})
    assert r.status_code == 400


@pytest.mark.parametrize(
    ("filename", "content_type", "content"),
    [
        ("fake.pdf", "application/pdf", b"not a pdf"),
        (
            "fake.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            b"not a zip archive",
        ),
        ("binary.txt", "text/plain", b"text\x00binary"),
    ],
)
def test_intake_document_rejects_disguised_or_binary_content(client, filename, content_type, content):
    r = client.post("/intake/document", files={"file": (filename, content, content_type)})
    assert r.status_code == 400


def test_intake_document_rejects_empty_upload(client):
    r = client.post("/intake/document", files={"file": ("empty.txt", b"", "text/plain")})
    assert r.status_code == 422


def test_intake_document_rejects_whitespace_only_txt(client):
    r = client.post("/intake/document", files={"file": ("blank.txt", b"  \n\t", "text/plain")})
    assert r.status_code == 422


def test_intake_document_enforces_upload_limit(client, monkeypatch):
    from app.config import settings

    monkeypatch.setattr(settings, "max_upload_bytes", 8)
    r = client.post("/intake/document", files={"file": ("large.txt", b"123456789", "text/plain")})
    assert r.status_code == 413


def test_encroachment_is_not_tagged_as_ncr():
    issues = [i["en"] for i in build_intake_result("Earthwork encroachment at the interface zone, Malaysia.")["issues"]]
    assert "Earthwork encroachment" in issues and "Non-conformance (NCR)" not in issues
    assert "Non-conformance (NCR)" in [i["en"] for i in build_intake_result("Raised NCR-2026-084 for honeycomb.")["issues"]]
